from fastapi import FastAPI, APIRouter, HTTPException, Body
from fastapi.responses import FileResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
import os
import re
import csv as csvlib
import logging
import uuid
import io
import smtplib
import json
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from pathlib import Path
from pydantic import BaseModel, EmailStr
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Optional


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

EXPORTS_DIR = ROOT_DIR / "exports"
EXPORTS_DIR.mkdir(exist_ok=True)

# ---------------------------------------------------------------------------
# Email (SMTP via Gmail) configuration
# ---------------------------------------------------------------------------
GMAIL_SENDER = os.environ.get("GMAIL_SENDER", "wajesanjana25@gmail.com")
GMAIL_APP_PASSWORD = os.environ.get("GMAIL_APP_PASSWORD")

app = FastAPI()
api_router = APIRouter(prefix="/api")


class AdminLoginRequest(BaseModel):
    email: str
    password: str


class AdminLoginResponse(BaseModel):
    success: bool
    token: Optional[str] = None


@api_router.get("/")
async def root():
    return {"message": "Feedback System API"}


@api_router.post("/admin/login", response_model=AdminLoginResponse)
async def admin_login(payload: AdminLoginRequest):
    expected_email = os.environ.get("ADMIN_EMAIL")
    expected_password = os.environ.get("ADMIN_PASSWORD")

    if payload.email.strip().lower() == expected_email.strip().lower() and payload.password == expected_password:
        return AdminLoginResponse(success=True, token="admin-session-ok")

    raise HTTPException(status_code=401, detail="Invalid email or password")


# ---------------------------------------------------------------------------
# Export columns (must match Supabase column names)
# ---------------------------------------------------------------------------

NUMERIC_QUESTION_COLS = [
    "how_would_you_rate_the_overall_content_of_the_session?",
    "were_the_topics_covered_relevant_to_your_needs_and_expectations",
    "rate_the_clarity_and_depth_of_the_information_provided",
    "how_knowledgeable_was_the_instructor_about the_subject_matter?",
    "how_engaging_was_the_instructor_in_delivering_the_content?",
    "rate_the_instructor's_ability_to_answer_questions_and_provide_e",
    "how_would_you_rate_the_overall_organization_and_flow_of_the_ses",
    "rate_the_activities/exercises_conducted_during_the_session_to_u",
    "rate_the_overall_interaction_and_engagement_opportunities_durin",
]

TEXT_QUESTION_COLS = [
    "what_did_you_like_most_about_the_session?",
    "what_areas_do_you_think_need_improvement",
]

# Full question-text labels (user-specified)
HEADER_LABELS = {
    "how_would_you_rate_the_overall_content_of_the_session?": "How would you rate the overall content of the session?",
    "were_the_topics_covered_relevant_to_your_needs_and_expectations": "Were the topics covered relevant to your needs and expectations?",
    "rate_the_clarity_and_depth_of_the_information_provided": "Rate the clarity and depth of the information provided.",
    "how_knowledgeable_was_the_instructor_about the_subject_matter?": "How knowledgeable was the instructor about the subject matter?",
    "how_engaging_was_the_instructor_in_delivering_the_content?": "How engaging was the instructor in delivering the content?",
    "rate_the_instructor's_ability_to_answer_questions_and_provide_e": "Rate the instructor's ability to answer questions and provide explanations.",
    "how_would_you_rate_the_overall_organization_and_flow_of_the_ses": "How would you rate the overall organization and flow of the session?",
    "rate_the_activities/exercises_conducted_during_the_session_to_u": "Rate the activities/exercises conducted during the session to understand the concepts?",
    "rate_the_overall_interaction_and_engagement_opportunities_durin": "Rate the overall interaction and engagement opportunities during the session.",
    "what_did_you_like_most_about_the_session?": "What did you like most about the session?",
    "what_areas_do_you_think_need_improvement": "What areas do you think need improvement",
}

META_COLS = ["date_time", "institution", "category", "program", "batch", "instructor"]
META_LABELS = ["Submitted Date and Time", "Institution", "Category", "Program", "Batch", "Instructor"]

# Editable, ops-only columns added from the Negative Feedback page UI.
# Always appended at the very end of the export (after email) when
# include_extra_cols is True, so the layout stays predictable.
EXTRA_COLS = ["allotted_to", "feedback"]
EXTRA_LABELS = ["Allotted To", "Feedback"]


class ExportRequest(BaseModel):
    responses: List[Dict[str, Any]]
    title: Optional[str] = None
    include_email: bool = True
    include_avg_row: bool = True
    include_extra_cols: bool = False  # only Negative Feedback page sets this True
    session_key: Optional[str] = None  # stable id (institution/category/program/batch/instructor/date)
                                        # -> when set, the export file is overwritten in place
                                        # instead of getting a brand-new random filename each time,
                                        # so a previously shared link always serves the latest data.


class SendEmailRequest(BaseModel):
    instructor_email: EmailStr
    instructor_name: str
    csv_url: str
    session_label: Optional[str] = ""


def _to_float(v):
    try:
        if v is None or v == "":
            return None
        return float(v)
    except (ValueError, TypeError):
        return None


def _sanitize_key(key: str) -> str:
    """Turns an arbitrary filter-derived string into a safe, stable filename
    stem: lowercased, non-alphanumerics collapsed to single underscores."""
    key = key.strip().lower()
    key = re.sub(r"[^a-z0-9]+", "_", key).strip("_")
    return key or "export"

def _sidecar_path(file_id: str) -> Path:
    return EXPORTS_DIR / f"{file_id}.json"


def _load_existing_responses(session_key: Optional[str]) -> List[Dict[str, Any]]:
    """Loads whatever raw rows were previously exported under this session_key,
    so a new fetch (e.g. a different date) can be merged in rather than
    replacing the file outright."""
    if not session_key:
        return []
    path = _sidecar_path(_sanitize_key(session_key))
    if not path.exists():
        return []
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        logger.exception(f"Failed to load sidecar data for session_key={session_key}")
        return []


def _row_identity(row: Dict[str, Any]):
    """Unique identity for a response row, used to de-dupe when merging.
    Prefers the Supabase primary key; falls back to a composite of
    fields that should uniquely identify a submission if 'id' is absent."""
    if row.get("id") is not None:
        return ("id", row.get("id"))
    return (
        "composite",
        row.get("date_time"),
        row.get("email_address"),
        row.get("institution"),
        row.get("program"),
        row.get("batch"),
        row.get("instructor"),
    )


def _merge_responses(existing: List[Dict[str, Any]], new: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Combines previously-exported rows with newly-fetched rows, de-duped
    by identity (a newer fetch of the same row wins), then sorts
    chronologically so sessions land in date order in the export."""
    merged: Dict[Any, Dict[str, Any]] = {}
    for row in existing:
        merged[_row_identity(row)] = row
    for row in new:
        merged[_row_identity(row)] = row
    combined = list(merged.values())
    combined.sort(key=lambda r: (r.get("date_time") or ""))
    return combined


def _save_sidecar(session_key: Optional[str], responses: List[Dict[str, Any]]):
    if not session_key:
        return
    path = _sidecar_path(_sanitize_key(session_key))
    try:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(responses, f)
    except Exception:
        logger.exception(f"Failed to save sidecar data for session_key={session_key}")
        
def _build_headers(include_email: bool, include_extra_cols: bool = False):
    headers = list(META_LABELS)
    headers += [HEADER_LABELS[c] for c in NUMERIC_QUESTION_COLS]
    headers += [HEADER_LABELS[c] for c in TEXT_QUESTION_COLS]
    if include_email:
        headers.append("Email Address")
    if include_extra_cols:
        headers += EXTRA_LABELS
    return headers


def _row_values(row: Dict[str, Any], include_email: bool, include_extra_cols: bool = False):
    meta = [row.get(c) for c in META_COLS]
    nums = [row.get(c) for c in NUMERIC_QUESTION_COLS]
    texts = [row.get(c) for c in TEXT_QUESTION_COLS]
    values = meta + nums + texts
    if include_email:
        values.append(row.get("email_address"))
    if include_extra_cols:
        values += [row.get(c, "") for c in EXTRA_COLS]
    return values


# ---------------------------------------------------------------------------
# Session grouping — mirrors the frontend's session grouping key
# (institution + category + program + batch + instructor + date) so that
# per-session average rows land in the right place in an export that spans
# multiple sessions (e.g. "all sessions of a program").
# ---------------------------------------------------------------------------

def _session_key(row: Dict[str, Any]):
    date_time = row.get("date_time") or ""
    date = date_time[:10] if date_time else "—"
    return (
        row.get("institution"),
        row.get("category"),
        row.get("program"),
        row.get("batch"),
        row.get("instructor"),
        date,
    )


def _group_by_session(responses: List[Dict[str, Any]]):
    """Groups rows by session, preserving first-seen order (dict keeps
    insertion order in Python 3.7+), and returns a list of (key, rows)."""
    groups: Dict[Any, List[Dict[str, Any]]] = {}
    for row in responses:
        key = _session_key(row)
        groups.setdefault(key, []).append(row)
    return list(groups.items())


def _avg_for_cols(rows_subset: List[Dict[str, Any]]):
    """Returns (list_of_column_averages, overall_average) for the numeric
    question columns across the given subset of rows."""
    column_averages = []
    for q_col in NUMERIC_QUESTION_COLS:
        vals = [_to_float(r.get(q_col)) for r in rows_subset]
        vals = [v for v in vals if v is not None]
        avg = round(sum(vals) / len(vals), 2) if vals else None
        column_averages.append(avg)
    valid = [a for a in column_averages if a is not None]
    overall = round(sum(valid) / len(valid), 2) if valid else None
    return column_averages, overall


def build_xlsx(
    responses: List[Dict[str, Any]],
    include_email: bool,
    include_avg_row: bool = True,
    include_extra_cols: bool = False,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Feedback"

    header_fill = PatternFill(start_color="DCEAD7", end_color="DCEAD7", fill_type="solid")
    avg_fill = PatternFill(start_color="9BE5F3", end_color="9BE5F3", fill_type="solid")
    overall_fill = PatternFill(start_color="F6B26B", end_color="F6B26B", fill_type="solid")
    bold = Font(bold=True)
    thin = Side(border_style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = _build_headers(include_email, include_extra_cols)
    numeric_start_col = len(META_LABELS) + 1  # 1-indexed first numeric column
    numeric_end_col = numeric_start_col + len(NUMERIC_QUESTION_COLS) - 1
    overall_col = numeric_end_col + 1

    # Header row
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = bold
        cell.alignment = center
        cell.border = border

    neg_row_fill = PatternFill(start_color="FFF8E1", end_color="FFF8E1", fill_type="solid")

    # Split into sessions so we can drop an average row after each one.
    # If include_avg_row is False, treat everything as a single "session"
    # so the loop below still just writes the raw rows.
    session_items = _group_by_session(responses) if include_avg_row else [(None, responses)]

    def _write_avg_row(row_idx, rows_subset, label, fill):
        cell = ws.cell(row=row_idx, column=1, value=label)
        cell.font = bold
        cell.fill = fill
        for c_i in range(2, numeric_start_col):
            ws.cell(row=row_idx, column=c_i).fill = fill
        column_averages, overall = _avg_for_cols(rows_subset)
        for q_i, avg in enumerate(column_averages):
            c_idx = numeric_start_col + q_i
            cell = ws.cell(row=row_idx, column=c_idx, value=avg)
            cell.fill = fill
            cell.font = bold
            cell.alignment = center
        cell = ws.cell(row=row_idx, column=overall_col, value=overall)
        cell.fill = overall_fill if fill is avg_fill else fill
        cell.font = bold
        cell.alignment = center

    current_row = 2
    for _key, group_rows in session_items:
        for row in group_rows:
            values = _row_values(row, include_email, include_extra_cols)

            neg_count = sum(
                1 for c in NUMERIC_QUESTION_COLS
                if _to_float(row.get(c)) is not None and _to_float(row.get(c)) <= 3
            )
            highlight_row = neg_count >= 3

            for c_i, v in enumerate(values, start=1):
                cell = ws.cell(row=current_row, column=c_i, value=v)
                cell.border = border
                cell.alignment = Alignment(horizontal='center', vertical="center", wrap_text=True)

                if highlight_row:
                    cell.fill = neg_row_fill

                if numeric_start_col <= c_i <= numeric_end_col:
                    fv = _to_float(v)
                    if fv is not None and fv <= 3:
                        cell.font = Font(bold=True, color="CC0000")
            current_row += 1

        if include_avg_row:
            _write_avg_row(current_row, group_rows, "Average", avg_fill)
            current_row += 1

    # Column widths
    widths = [22, 26, 14, 14, 16, 24]
    widths += [16] * len(NUMERIC_QUESTION_COLS)
    widths += [40, 40]
    if include_email:
        widths.append(28)
    if include_extra_cols:
        widths += [20, 32]  # Allotted To, Feedback
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 70

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(
    responses: List[Dict[str, Any]],
    include_email: bool,
    include_avg_row: bool = True,
    include_extra_cols: bool = False,
) -> bytes:
    headers = _build_headers(include_email, include_extra_cols)
    buf = io.StringIO()
    writer = csvlib.writer(buf)
    writer.writerow(headers)

    def _avg_row_for(rows_subset, label):
        avg_row = [label] + [""] * (len(META_LABELS) - 1)
        column_averages, overall = _avg_for_cols(rows_subset)
        avg_row += column_averages
        avg_row.append(overall)  # overall column sits right after the numeric block
        avg_row.append("")  # second text-question column stays blank
        if include_email:
            avg_row.append("")
        if include_extra_cols:
            avg_row += [""] * len(EXTRA_COLS)
        return avg_row

    session_items = _group_by_session(responses) if include_avg_row else [(None, responses)]

    for _key, group_rows in session_items:
        for row in group_rows:
            writer.writerow(_row_values(row, include_email, include_extra_cols))
        if include_avg_row:
            writer.writerow(_avg_row_for(group_rows, "Average"))

    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 nicely


MIME_TYPES = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
}


def _save_and_url(
    data: bytes,
    ext: str,
    basename: str = "feedback",
    session_key: Optional[str] = None,
) -> Dict[str, str]:
    # If a session_key is supplied, derive a stable, sanitized file_id from it
    # so re-exporting the same filter combo overwrites the same file instead
    # of creating a new one — this is what makes a previously shared link
    # ("Send to instructor") automatically reflect newer data on a re-fetch.
    # Without a session_key (e.g. ad-hoc/negative-feedback exports), fall
    # back to the old random-UUID behavior so those stay one-off snapshots.
    if session_key:
        file_id = _sanitize_key(session_key)
    else:
        file_id = uuid.uuid4().hex

    path = EXPORTS_DIR / f"{file_id}.{ext}"
    with open(path, "wb") as f:
        f.write(data)

    backend_base = os.environ.get("PUBLIC_BACKEND_URL", "").rstrip("/")
    relative = f"/api/export/file/{file_id}.{ext}"
    url = f"{backend_base}{relative}" if backend_base else relative

    return {"file_id": file_id, "url": url}


@api_router.post("/export/xlsx")
async def export_xlsx(payload: ExportRequest = Body(...)):
    if not payload.responses:
        raise HTTPException(status_code=400, detail="No responses to export")
    existing = _load_existing_responses(payload.session_key)
    combined = _merge_responses(existing, payload.responses) if payload.session_key else payload.responses
    data = build_xlsx(
        combined,
        payload.include_email,
        payload.include_avg_row,
        payload.include_extra_cols,
    )
    _save_sidecar(payload.session_key, combined)
    return _save_and_url(data, "xlsx", basename="feedback_session", session_key=payload.session_key)


@api_router.post("/export/csv")
async def export_csv(payload: ExportRequest = Body(...)):
    if not payload.responses:
        raise HTTPException(status_code=400, detail="No responses to export")
    existing = _load_existing_responses(payload.session_key)
    combined = _merge_responses(existing, payload.responses) if payload.session_key else payload.responses
    data = build_csv(
        combined,
        payload.include_email,
        payload.include_avg_row,
        payload.include_extra_cols,
    )
    _save_sidecar(payload.session_key, combined)
    return _save_and_url(data, "csv", basename="feedback_combined", session_key=payload.session_key)

@api_router.post("/export/send-email")
async def send_email(payload: SendEmailRequest = Body(...)):
    """Emails a previously-generated CSV/XLSX export link to an instructor
    via Gmail SMTP. Requires GMAIL_APP_PASSWORD to be set in the backend
    environment (a Gmail App Password for GMAIL_SENDER, not the normal
    account password)."""
    if not GMAIL_APP_PASSWORD:
        raise HTTPException(
            status_code=500,
            detail="Email sending is not configured (missing GMAIL_APP_PASSWORD)",
        )

    subject = f"Feedback Report{' – ' + payload.session_label if payload.session_label else ''}"
    body = f"""Hi {payload.instructor_name},

Your session feedback report is ready. You can view/download it here:

{payload.csv_url}

Best,
Accredian Team
"""

    msg = MIMEMultipart()
    msg["From"] = GMAIL_SENDER
    msg["To"] = payload.instructor_email
    msg["Subject"] = subject
    msg.attach(MIMEText(body, "plain"))

    try:
        with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
            server.login(GMAIL_SENDER, GMAIL_APP_PASSWORD)
            server.send_message(msg)
    except smtplib.SMTPAuthenticationError:
        logger.error("Gmail SMTP authentication failed — check GMAIL_SENDER / GMAIL_APP_PASSWORD")
        raise HTTPException(
            status_code=500,
            detail="Email authentication failed. Check GMAIL_SENDER and GMAIL_APP_PASSWORD.",
        )
    except Exception as e:
        logger.error(f"Email send failed: {e}")
        raise HTTPException(status_code=500, detail=f"Email send failed: {str(e)}")

    return {"status": "sent", "to": payload.instructor_email}


# @api_router.post("/export/negative-xlsx")
# async def export_negative_xlsx(payload: ExportRequest = Body(...)):
#     if not payload.responses:
#         raise HTTPException(status_code=400, detail="No responses to export")
#     data = build_xlsx(payload.responses, payload.include_email, include_avg_row=False)
#     return _save_and_url(data, "xlsx", basename="negative_feedback")


# @api_router.post("/export/negative-csv")
# async def export_negative_csv(payload: ExportRequest = Body(...)):
#     if not payload.responses:
#         raise HTTPException(status_code=400, detail="No responses to export")
#     data = build_csv(payload.responses, payload.include_email, include_avg_row=False)
#     return _save_and_url(data, "csv", basename="negative_feedback")


@api_router.get("/export/file/{filename}")
async def serve_export(filename: str):
    if "/" in filename or ".." in filename:
        raise HTTPException(status_code=400, detail="Invalid filename")
    path = EXPORTS_DIR / filename
    if not path.exists():
        raise HTTPException(status_code=404, detail="Export not found")

    if filename.endswith(".xlsx"):
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif filename.endswith(".csv"):
        media = "text/csv"
    else:
        media = "application/octet-stream"

    return FileResponse(path, media_type=media, filename=f"feedback_{filename}")


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)